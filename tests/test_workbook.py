import csv
import shutil
import subprocess

import pytest
from build_workbook import (
    CHECKLIST,
    EXAMPLE_ROWS,
    OVERRIDE_LABEL,
    REPLACE_LABEL,
    build_workbook,
    checklist_markdown,
    score,
)
from openpyxl import load_workbook

CSV_FILTER = "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1"
needs_soffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="no LibreOffice")


def _recalculate(xlsx, out_dir):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", CSV_FILTER, "--outdir", str(out_dir), str(xlsx)],
        check=True, capture_output=True, timeout=180,
    )


def _read(path):
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.reader(handle))


def test_score_matches_formula():
    assert score(5, 0.8, 2, 2) == 1.0


def test_zero_divisor_returns_none():
    assert score(3, 0.5, 0, 0) is None


@pytest.mark.parametrize("confidence", [80, 0, 0.05, 1.5])
def test_confidence_must_be_a_fraction(confidence):
    with pytest.raises(ValueError):
        score(3, confidence, 2, 2)


@pytest.mark.parametrize("value", [0, 6])
def test_value_must_be_1_to_5(value):
    with pytest.raises(ValueError):
        score(value, 0.5, 2, 2)


def test_workbook_has_live_formulas(tmp_path):
    path = tmp_path / "w.xlsx"
    build_workbook(path, EXAMPLE_ROWS)
    ws = load_workbook(path)["Inventory"]

    assert ws["L2"].value.startswith("=IF(")
    assert "CHECK INPUT" in ws["L2"].value
    assert ws["L150"].value.startswith("=IF(")  # rows added later are covered


def test_build_rejects_bad_input_before_writing(tmp_path):
    path = tmp_path / "w.xlsx"
    with pytest.raises(ValueError):
        build_workbook(path, [dict(EXAMPLE_ROWS[0], confidence=80)])
    assert not path.exists()


def test_checklist_covers_governance_and_markdown_matches_sheet():
    md = checklist_markdown()
    for item_id, question, _, _ in CHECKLIST:
        assert item_id in md and question in md
    text = " ".join(" ".join(c) for c in CHECKLIST).lower()
    for needle in ("mfa", "least-privilege", "approved provider", "fp&a", "read-only", "gdpr"):
        assert needle in text


@needs_soffice
def test_recalculated_ranking_and_decisions(tmp_path):
    path = tmp_path / "w.xlsx"
    build_workbook(path, EXAMPLE_ROWS)
    _recalculate(path, tmp_path)
    ranking = _read(tmp_path / "w-Ranking.csv")
    body = ranking[1 : 1 + len(EXAMPLE_ROWS)]

    expected = sorted(
        EXAMPLE_ROWS, key=lambda r: -score(r["value"], r["confidence"], r["effort"], r["risk"])
    )
    assert [r[1] for r in body] == [r["tool"] for r in expected]
    scores = [float(r[2]) for r in body]
    assert scores == sorted(scores, reverse=True)
    decisions = {r[1]: r[3] for r in body}
    assert decisions["Site monitoring portal (vendor)"] == OVERRIDE_LABEL  # 1.25 but critical
    assert decisions["Weekly KPI report (hand-built in spreadsheets)"] == REPLACE_LABEL
    assert decisions["Accounting software"] == OVERRIDE_LABEL
    # only the forms SaaS has a cost among replace candidates (the report costs nothing)
    assert any(row[6:8] == ["Spend on replace candidates (USD)", "120"] for row in ranking)


@needs_soffice
def test_ties_rank_in_row_order(tmp_path):
    rows = [
        dict(EXAMPLE_ROWS[0], tool="First"),
        dict(EXAMPLE_ROWS[0], tool="Second"),
    ]
    path = tmp_path / "w.xlsx"
    build_workbook(path, rows)
    _recalculate(path, tmp_path)
    ranking = _read(tmp_path / "w-Ranking.csv")

    assert [ranking[1][1], ranking[2][1]] == ["First", "Second"]


@needs_soffice
def test_rows_appended_below_the_table_are_scored_and_ranked(tmp_path):
    path = tmp_path / "w.xlsx"
    build_workbook(path, EXAMPLE_ROWS[:2])
    wb = load_workbook(path)
    ws = wb["Inventory"]
    for col, val in zip("ABGHIJK", ("Added later", "Test", "N", 5, 1.0, 1, 1), strict=True):
        ws[f"{col}4"] = val  # row 4 is empty; a user would type here
    wb.save(path)
    _recalculate(path, tmp_path)
    ranking = _read(tmp_path / "w-Ranking.csv")

    assert ranking[1][1] == "Added later"  # 5 * 1.0 / 2 = 2.5 beats the others


@needs_soffice
def test_pasted_zero_effort_and_risk_shows_check_input(tmp_path):
    path = tmp_path / "w.xlsx"
    build_workbook(path, [dict(EXAMPLE_ROWS[0])])
    wb = load_workbook(path)
    wb["Inventory"]["J2"], wb["Inventory"]["K2"] = 0, 0  # bypasses validation, as a paste would
    wb.save(path)
    _recalculate(path, tmp_path)
    inventory = _read(tmp_path / "w-Inventory.csv")

    assert inventory[1][11] == "CHECK INPUT"
    assert inventory[1][12] == "CHECK INPUT"


def test_generated_pages_are_up_to_date():
    from pathlib import Path

    from build_workbook import worked_example_markdown

    root = Path(__file__).resolve().parent.parent
    assert (root / "worked-example.md").read_text(encoding="utf-8") == worked_example_markdown()
    assert (root / "decision-checklist.md").read_text(encoding="utf-8") == checklist_markdown()


@needs_soffice
def test_python_decisions_agree_with_sheet_formulas_for_every_row(tmp_path):
    from build_workbook import decide

    path = tmp_path / "w.xlsx"
    build_workbook(path, EXAMPLE_ROWS)
    _recalculate(path, tmp_path)
    sheet = {r[1]: r[3] for r in _read(tmp_path / "w-Ranking.csv")[1 : 1 + len(EXAMPLE_ROWS)]}

    assert sheet == {r["tool"]: decide(r) for r in EXAMPLE_ROWS}
